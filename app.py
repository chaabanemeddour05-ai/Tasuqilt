
import io, os, re, json, time, unicodedata
from difflib import SequenceMatcher
from datetime import datetime, timezone
import requests, pandas as pd, streamlit as st
from docx import Document
from openpyxl import load_workbook

APP_TITLE = "Tasuqilt DZ"
DEFAULT_GEMINI_MODEL = "gemini-3.1-flash-lite"
TIMEOUT = 30
DIRECTIONS = {
    "Français → Tamazight": ("fr", "tz"),
    "العربية → Tamazight": ("ar", "tz"),
    "Tamazight → Français": ("tz", "fr"),
    "Tamazight → العربية": ("tz", "ar"),
}
DEFAULT_TERMS = [
    {"source":"Président de la République", "target":"Aselway n Tegduda", "source_lang":"fr", "target_lang":"tz", "status":"approved", "origin":"Tasuqilt core"},
    {"source":"Président", "target":"Aselway", "source_lang":"fr", "target_lang":"tz", "status":"approved", "origin":"Tasuqilt core"},
    {"source":"Conseil des ministres", "target":"Aseqqamu n Yineɣlafen", "source_lang":"fr", "target_lang":"tz", "status":"approved", "origin":"Tasuqilt core"},
    {"source":"Réunion", "target":"Timlilt", "source_lang":"fr", "target_lang":"tz", "status":"approved", "origin":"Tasuqilt core"},
    {"source":"Alger", "target":"DZAYER TAMANEƔT", "source_lang":"fr", "target_lang":"tz", "status":"approved", "origin":"Tasuqilt core"},
    {"source":"Gouvernement", "target":"Anabaḍ", "source_lang":"fr", "target_lang":"tz", "status":"approved", "origin":"Tasuqilt core"},
]

# Settings are read from Streamlit Secrets first, then environment variables.
def setting(name, default=""):
    try:
        value = st.secrets.get(name, default)
        if value not in (None, ""):
            return str(value).strip()
    except Exception:
        pass
    return str(os.environ.get(name, default)).strip()

GEMINI_API_KEY = setting("GEMINI_API_KEY")
GEMINI_MODEL = setting("GEMINI_MODEL", DEFAULT_GEMINI_MODEL)
SUPABASE_URL = setting("SUPABASE_URL").rstrip("/")
SUPABASE_KEY = setting("SUPABASE_SERVICE_ROLE_KEY")
ADMIN_PASSWORD = setting("ADMIN_PASSWORD")
ASEGZAWAL_API = setting("ASEGZAWAL_API", "https://asegzawal.miraheze.org/w/api.php")


def norm(value):
    if value is None: return ""
    value = unicodedata.normalize("NFKC", str(value)).replace("\u00a0", " ").replace("\u200b", "")
    return re.sub(r"\s+", " ", value).strip().casefold()


def norm_search(value):
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", norm(value), flags=re.UNICODE)).strip()


def clean(value):
    return "" if value is None else str(value).strip()


def similarity(a, b):
    aa, bb = norm_search(a), norm_search(b)
    if not aa or not bb: return 0.0
    ta, tb = set(aa.split()), set(bb.split())
    jac = len(ta & tb) / max(1, len(ta | tb))
    return max(jac, SequenceMatcher(None, aa, bb).ratio() * 0.72)


def db_ready():
    return bool(SUPABASE_URL and SUPABASE_KEY)


def db_headers(prefer=None):
    h = {
        "apikey": SUPABASE_KEY,
        "Content-Type": "application/json",
    }
    if prefer:
        h["Prefer"] = prefer
    return h


```python
def db_select(table, params=None, limit=3000):
    if not db_ready():
        return []

    q = dict(params or {})
    q.setdefault("select", "*")

    try:
        total_limit = int(limit)
    except (TypeError, ValueError):
        total_limit = 3000

    if total_limit <= 0:
        return []

    page_size = 1000
    results = []
    offset = 0

    while len(results) < total_limit:
        current_limit = min(page_size, total_limit - len(results))
        page_params = dict(q)
        page_params["limit"] = str(current_limit)

        headers = db_headers()
        headers["Range-Unit"] = "items"
        headers["Range"] = f"{offset}-{offset + current_limit - 1}"

        response = requests.get(
            f"{SUPABASE_URL}/rest/v1/{table}",
            headers=headers,
            params=page_params,
            timeout=TIMEOUT,
        )
        response.raise_for_status()

        data = response.json()

        if not isinstance(data, list) or not data:
            break

        results.extend(data)
        offset += len(data)

        if len(data) < current_limit:
            break

    return results
```


def db_insert(table, record):
    if not db_ready(): raise RuntimeError("Supabase غير مضبوط؛ لا يمكن حفظ البيانات بصورة دائمة.")
    r = requests.post(f"{SUPABASE_URL}/rest/v1/{table}", headers=db_headers("return=representation"), json=record, timeout=TIMEOUT)
    r.raise_for_status(); data = r.json()
    return data[0] if isinstance(data, list) and data else data


def db_upsert(table, records):
    if not db_ready(): raise RuntimeError("Supabase غير مضبوط.")
    if not records: return
    h = db_headers("resolution=ignore-duplicates,return=minimal")
    for start in range(0, len(records), 100):
        r = requests.post(f"{SUPABASE_URL}/rest/v1/{table}", headers=h, json=records[start:start+100], timeout=TIMEOUT)
        r.raise_for_status()


def now_iso(): return datetime.now(timezone.utc).isoformat()


def make_pair(source, target, origin="", source_lang="fr", target_lang="tz"):
    source, target = clean(source), clean(target)
    if not source or not target or norm(source) == norm(target): return None
    return {"source":source, "target":target, "origin":origin, "source_lang":source_lang, "target_lang":target_lang}


def parse_file(filename, content):
    ext = os.path.splitext(filename)[1].lower(); out = []
    if ext == ".xlsx":
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = next((s for s in wb.worksheets if s.title.strip().casefold() == "tm"), wb.worksheets[0])
            for i, row in enumerate(sheet.iter_rows(min_col=1, max_col=2, values_only=True)):
                if not row or len(row) < 2: continue
                a, b = clean(row[0]), clean(row[1])
                if i == 0 and norm(a) in {"source", "français", "francais", "original", "العربية"}: continue
                p = make_pair(a, b, filename)
                if p: out.append(p)
        finally: wb.close()
    elif ext == ".csv":
        try: df = pd.read_csv(io.BytesIO(content), dtype=str, keep_default_na=False, encoding="utf-8-sig")
        except Exception: df = pd.read_csv(io.BytesIO(content), dtype=str, keep_default_na=False, encoding="latin-1")
        if len(df.columns) >= 2:
            for i, row in df.iterrows():
                a, b = clean(row.iloc[0]), clean(row.iloc[1])
                if i == 0 and norm(a) in {"source", "français", "francais", "original", "العربية"}: continue
                p = make_pair(a, b, filename)
                if p: out.append(p)
    elif ext == ".json":
        data = json.loads(content.decode("utf-8-sig"))
        if isinstance(data, dict): data = data.get("pairs", data.get("translations", [data]))
        if isinstance(data, list):
            for item in data:
                if isinstance(item, dict):
                    p = make_pair(item.get("source", item.get("fr", item.get("original", ""))), item.get("target", item.get("tz", item.get("tamazight", item.get("translation", "")))), filename)
                    if p: out.append(p)
    elif ext == ".docx":
        doc = Document(io.BytesIO(content))
        for para in doc.paragraphs:
            if "@" in para.text:
                a, b = para.text.split("@", 1); p = make_pair(a, b, filename)
                if p: out.append(p)
        for table in doc.tables:
            for row in table.rows:
                if len(row.cells) >= 2:
                    p = make_pair(row.cells[0].text, row.cells[1].text, filename)
                    if p: out.append(p)
    elif ext in {".txt", ".md"}:

        for line in content.decode("utf-8-sig", errors="replace").splitlines():
            if "@" in line:
                a, b = line.split("@", 1)
                p = make_pair(a, b, filename)
                if p:
                    out.append(p)
    return out


# ملفات GitHub مصدر للقراءة فقط؛ التعديلات الجديدة تحفظ في Supabase.
@st.cache_data(ttl=180, show_spinner=False)
def load_repo_pairs():
    owner, repo = "chaabanemeddour05-ai", "Tasuqilt"
    api = f"https://api.github.com/repos/{owner}/{repo}"
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "Tasuqilt-DZ",
    }
    token = setting("GITHUB_TOKEN")
    if token:
        headers["Authorization"] = f"Bearer {token}"

    try:
        info = requests.get(api, headers=headers, timeout=TIMEOUT)
        info.raise_for_status()
        branch = info.json().get("default_branch", "main")

        tree_url = f"{api}/git/trees/{branch}?recursive=1"
        tree_r = requests.get(tree_url, headers=headers, timeout=TIMEOUT)
        tree_r.raise_for_status()
        tree = tree_r.json()

        if tree.get("truncated"):
            raise RuntimeError("GitHub أعاد قائمة ملفات غير مكتملة.")

        pairs, errors = [], []
        allowed = {".xlsx", ".csv", ".json", ".txt", ".md", ".docx"}

        for item in tree.get("tree", []):
            path = item.get("path", "")
            ext = os.path.splitext(path)[1].lower()

            if (
                item.get("type") != "blob"
                or "/" in path
                or ext not in allowed
                or item.get("size", 0) > 20 * 1024 * 1024
            ):
                continue

            url = (
                f"https://raw.githubusercontent.com/"
                f"{owner}/{repo}/{branch}/"
                f"{requests.utils.quote(path)}"
            )

            try:
                r = requests.get(
                    url,
                    headers={"User-Agent": "Tasuqilt-DZ"},
                    timeout=TIMEOUT,
                )
                r.raise_for_status()
                pairs.extend(parse_file(path, r.content))
            except Exception as exc:
                errors.append(f"{path}: {exc}")

        unique = {
            (norm(p["source"]), norm(p["target"])): p
            for p in pairs
        }
        return list(unique.values()), errors

    except Exception as exc:
        return [], [f"تعذر تحميل ملفات GitHub: {exc}"]


def get_terms():
    terms = [dict(t) for t in DEFAULT_TERMS]

    if db_ready():
        try:
            rows = db_select(
                "lexicon_entries",
                {
                    "status": "eq.approved",
                    "order": "created_at.desc",
                },
                5000,
            )

            for row in rows:
                key = (
                    norm(row.get("source")),
                    row.get("source_lang", "fr"),
                    row.get("target_lang", "tz"),
                )

                terms = [
                    t for t in terms
                    if (
                        norm(t.get("source")),
                        t.get("source_lang", "fr"),
                        t.get("target_lang", "tz"),
                    ) != key
                ]
                terms.insert(0, row)

        except Exception:
            pass

    return terms


def get_pairs(repo_pairs):
    merged = {}

    for p in repo_pairs:
        key = (
            norm(p.get("source")),
            norm(p.get("target")),
            p.get("source_lang", "fr"),
            p.get("target_lang", "tz"),
        )
        merged[key] = p

    if db_ready():
        try:
            rows = db_select(
                "translation_memory",
                {
                    "status": "eq.approved",
                    "order": "created_at.desc",
                },
                5000,
            )

            for p in rows:
                key = (
                    norm(p.get("source")),
                    norm(p.get("target")),
                    p.get("source_lang", "fr"),
                    p.get("target_lang", "tz"),
                )
                merged[key] = p

        except Exception:
            pass

    return list(merged.values())


def exact_lookup(text, direction, terms, pairs):
    src, tgt = DIRECTIONS[direction]
    q = norm(text)

    for item in terms + pairs:
        sl = item.get("source_lang", "fr")
        tl = item.get("target_lang", "tz")

        if sl == src and tl == tgt and norm(item.get("source")) == q:
            kind = (
                "مصطلح معتمد"
                if item in terms
                else "ذاكرة ترجمة معتمدة"
            )
            return item.get("target"), kind

        if sl == tgt and tl == src and norm(item.get("target")) == q:
            return item.get("source"), "بحث عكسي في المعرفة المعتمدة"

    return None, None


def relevant_terms(text, direction, terms, limit=60):
    src, tgt = DIRECTIONS[direction]
    q = norm_search(text)
    found = []

    for term in terms:
        if (
            term.get("source_lang", "fr") != src
            or term.get("target_lang", "tz") != tgt
        ):
            continue

        phrase = norm_search(term.get("source", ""))

        if phrase and (
            phrase in q
            or (
                len(phrase.split()) == 1
                and phrase in q.split()
            )
        ):
            found.append(term)

    found.sort(
        key=lambda x: len(norm_search(x.get("source", ""))),
        reverse=True,
    )
    return found[:limit]


def retrieve_examples(text, direction, pairs, limit=6):
    src, tgt = DIRECTIONS[direction]
    scored = []

    for p in pairs:
        sl = p.get("source_lang", "fr")
        tl = p.get("target_lang", "tz")

        if (sl, tl) == (src, tgt):
            a, b = p.get("source", ""), p.get("target", "")
        elif (sl, tl) == (tgt, src):
            a, b = p.get("target", ""), p.get("source", "")
        else:
            continue

        score = similarity(text, a)

        if score >= 0.13:
            scored.append({
                "source": a,
                "target": b,
                "score": score,
                "origin": p.get("origin", ""),
            })

    scored.sort(key=lambda x: x["score"], reverse=True)
    return scored[:limit]


def get_corpus():
    if not db_ready():
        return []

    try:
        return db_select(
            "language_corpus",
            {
                "language": "eq.tz",
                "status": "eq.approved",
                "select": "id,text,language,origin,status",
            },
            2000,
        )
    except Exception:
        return []


def retrieve_corpus_examples(text, corpus, limit=3):
    scored = []

    for item in corpus:
        paragraph = item.get("text", "")
        score = similarity(text, paragraph)

        if score >= 0.16:
            scored.append({
                "text": paragraph,
                "score": score,
                "origin": item.get("origin", ""),
            })

    scored.sort(key=lambda x: x["score"], reverse=True)
    return scored[:limit]

def gemini_translate(text, direction, terms, pairs, corpus):
    if not GEMINI_API_KEY:
        raise RuntimeError(
            "مفتاح Gemini غير مضبوط. أضف GEMINI_API_KEY إلى Streamlit Secrets."
        )

    src, tgt = DIRECTIONS[direction]
    term_hits = relevant_terms(text, direction, terms)
    examples = retrieve_examples(text, direction, pairs)
    corpus_hits = retrieve_corpus_examples(text, corpus) if src == "tz" else []

    lang_names = {
        "fr": "French",
        "ar": "Modern Standard Arabic",
        "tz": "Standard Amazigh (Tamazight)",
    }

    term_text = "\n".join(
        f"- {x.get('source', '')} → {x.get('target', '')}"
        for x in term_hits
    ) or "(No matching approved terminology)"

    example_text = "\n".join(
        f"- Source: {x['source']}\n  Approved translation: {x['target']}"
        for x in examples
    ) or "(No matching translation-memory examples)"

    corpus_text = "\n".join(
        f"- {x['text']}" for x in corpus_hits
    ) or "(No relevant Amazigh corpus paragraphs)"

    prompt = f"""
You are the translation engine of Tasuqilt DZ, a rigorous
standard Amazigh translation platform for Algerian public media.

Translate the complete input faithfully and naturally.

SOURCE LANGUAGE: {lang_names[src]}
TARGET LANGUAGE: {lang_names[tgt]}

NON-NEGOTIABLE RULES:
1. Preserve all facts, names, numbers, dates, quotations, and meaning.
2. Write in standard, consistent Amazigh when the target is Amazigh.
3. Approved terminology and approved translation-memory examples
   are authoritative guidance. Follow them in context.
4. Do not replace words mechanically when grammar or morphology
   requires a different form.
5. The approved lexical base form of "President" is Aselway.
   Do not use Anmazul as a translation of President.
6. The approved expression "President of the Republic" is
   Aselway n Tegduda.
7. The approved translation of "Réunion" is Timlilt.
8. The approved translation of "Conseil des ministres" is
   Aseqqamu n Yineɣlafen.
9. The approved translation of "Gouvernement" is Anabaḍ.
10. The approved rendering of "Alger" is DZAYER TAMANEƔT.
11. Respect Amazigh grammar, agreement, word order, and morphology.
12. Do not invent facts or add explanations.
13. Return only the translation, with no introduction or commentary.
14. If the source is already in the target language, revise it only
    when necessary for correctness and standardization.

APPROVED TERMINOLOGY:
{term_text}

APPROVED TRANSLATION-MEMORY EXAMPLES:
{example_text}

RELEVANT AMAZIGH CORPUS EXAMPLES:
{corpus_text}

TEXT TO TRANSLATE:
{text}
"""

    endpoint = (
        "https://generativelanguage.googleapis.com/v1beta/models/"
        f"{GEMINI_MODEL}:generateContent"
    )

    last_error = None

    for attempt in range(3):
        try:
            response = requests.post(
                endpoint,
                params={"key": GEMINI_API_KEY},
                json={
                    "contents": [{
                        "role": "user",
                        "parts": [{"text": prompt}],
                    }],
                    "generationConfig": {
                        "temperature": 0.15,
                        "maxOutputTokens": 4096,
                    },
                },
                timeout=90,
            )

            if response.status_code in (429, 500, 502, 503, 504):
                last_error = (
                    f"Gemini HTTP {response.status_code}: "
                    f"{response.text[:500]}"
                )
                if attempt < 2:
                    time.sleep(1.5 * (attempt + 1))
                    continue

            response.raise_for_status()
            data = response.json()

            candidates = data.get("candidates", [])
            if not candidates:
                raise RuntimeError(
                    "لم يُرجع Gemini ترجمة. قد يكون الطلب محجوبًا "
                    "أو لم ينتج النموذج أي نص."
                )

            parts = candidates[0].get("content", {}).get("parts", [])
            answer = "\n".join(
                part.get("text", "")
                for part in parts
                if part.get("text")
            ).strip()

            if not answer:
                raise RuntimeError("الترجمة التي أعادها Gemini فارغة.")

            return {
                "translation": answer,
                "terms": term_hits,
                "examples": examples,
                "corpus_examples": corpus_hits,
                "model": GEMINI_MODEL,
            }

        except requests.RequestException as exc:
            last_error = str(exc)
            if attempt < 2:
                time.sleep(1.5 * (attempt + 1))
                continue
            break

    raise RuntimeError(
        f"تعذرت الترجمة بعد عدة محاولات. التفاصيل: {last_error}"
    )


def search_asegzawal(query, limit=8):
    query = clean(query)

    if not query:
        return []

    params = {
        "action": "query",
        "format": "json",
        "list": "search",
        "srsearch": query,
        "srnamespace": 0,
        "srlimit": limit,
    }

    try:
        response = requests.get(
            ASEGZAWAL_API,
            params=params,
            timeout=TIMEOUT,
            headers={"User-Agent": "Tasuqilt-DZ/1.0"},
        )
        response.raise_for_status()
        data = response.json()

        results = []
        for item in data.get("query", {}).get("search", []):
            title = item.get("title", "")
            page_url = (
                "https://asegzawal.miraheze.org/wiki/"
                + requests.utils.quote(title.replace(" ", "_"))
            )
            results.append({
                "title": title,
                "snippet": re.sub(
                    r"<[^>]+>", "", item.get("snippet", "")
                ),
                "url": page_url,
            })

        return results

    except Exception:
        return []


def is_admin():
    return bool(
        ADMIN_PASSWORD
        and st.session_state.get("admin_authenticated", False)
    )


def require_admin():
    if not is_admin():
        st.error("هذه العملية متاحة للمشرفين المسجلين فقط.")
        st.stop()


def save_approved_pair(source, target, source_lang, target_lang, origin="admin"):
    require_admin()

    pair = make_pair(
        source,
        target,
        origin,
        source_lang,
        target_lang,
    )

    if not pair:
        raise ValueError("أدخل نص المصدر والترجمة بصورة صحيحة.")

    pair["status"] = "approved"
    pair["created_at"] = now_iso()

    db_insert("translation_memory", pair)
    return pair


def save_approved_term(source, target, source_lang, target_lang, note=""):
    require_admin()

    source, target = clean(source), clean(target)

    if not source or not target:
        raise ValueError("يجب إدخال المصطلح المصدر والمصطلح المعتمد.")

    record = {
        "source": source,
        "target": target,
        "source_lang": source_lang,
        "target_lang": target_lang,
        "status": "approved",
        "note": clean(note),
        "created_at": now_iso(),
    }

    db_insert("lexicon_entries", record)
    return record


def save_correction(source, proposed, corrected, direction, origin="Gemini"):
    require_admin()

    src, tgt = DIRECTIONS[direction]
    source, proposed, corrected = (
        clean(source),
        clean(proposed),
        clean(corrected),
    )

    if not source or not corrected:
        raise ValueError("النص الأصلي والتصحيح المعتمد مطلوبان.")

    if norm(source) == norm(corrected):
        raise ValueError("النص الأصلي مطابق للتصحيح؛ راجع المدخلات.")

    pair = {
        "source": source,
        "target": corrected,
        "source_lang": src,
        "target_lang": tgt,
        "origin": "approved correction",
        "status": "approved",
        "created_at": now_iso(),
    }

    db_insert("translation_memory", pair)

    correction = {
        "source": source,
        "proposed_translation": proposed,
        "corrected_translation": corrected,
        "source_lang": src,
        "target_lang": tgt,
        "origin": origin,
        "status": "approved",
        "created_at": now_iso(),
    }

    db_insert("corrections", correction)
    return correction


def save_corpus_paragraphs(paragraphs, origin="admin import"):
    require_admin()

    records = []
    seen = set()

    for paragraph in paragraphs:
        paragraph = clean(paragraph)

        if len(paragraph) < 20:
            continue

        key = norm(paragraph)
        if key in seen:
            continue

        seen.add(key)
        records.append({
            "text": paragraph,
            "language": "tz",
            "origin": origin,
            "status": "approved",
            "created_at": now_iso(),
        })

    if not records:
        return 0

    db_upsert("language_corpus", records)
    return len(records)

# ============================================================
# Tasuqilt DZ — Main Streamlit interface
# ============================================================

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="🌐",
    layout="wide",
)

st.title("Tasuqilt DZ")
st.caption(
    "منصة الترجمة الأمازيغية المعيارية — "
    "ذاكرة ترجمة، مصطلحات معتمدة، ومساعدة بالذكاء الاصطناعي"
)

with st.sidebar:
    st.header("إعدادات المنصة")
    st.write("**اتجاهات الترجمة**")
    st.write("الفرنسية ↔ الأمازيغية")
    st.write("العربية ↔ الأمازيغية")

    if db_ready():
        st.success("قاعدة البيانات: متصلة بالإعدادات")
    else:
        st.warning("قاعدة البيانات الدائمة غير مهيأة")

    if GEMINI_API_KEY:
        st.success("مفتاح Gemini موجود")
    else:
        st.warning("مفتاح Gemini غير مضبوط")

    st.divider()
    st.markdown(
        "[معجم Asegzawal](https://asegzawal.miraheze.org/wiki/Main_Page)"
    )

    with st.expander("دخول المشرف"):
        if not ADMIN_PASSWORD:
            st.warning(
                "يجب ضبط ADMIN_PASSWORD في Streamlit Secrets."
            )
        elif not is_admin():
            password = st.text_input(
                "كلمة مرور المشرف",
                type="password",
                key="admin_password_input",
            )
            if st.button("تسجيل الدخول", key="admin_login"):
                if password == ADMIN_PASSWORD:
                    st.session_state["admin_authenticated"] = True
                    st.success("تم تسجيل الدخول.")
                    st.rerun()
                else:
                    st.error("كلمة المرور غير صحيحة.")
        else:
            st.success("أنت مسجل الدخول بصفة مشرف.")
            if st.button("تسجيل الخروج", key="admin_logout"):
                st.session_state["admin_authenticated"] = False
                st.rerun()


@st.cache_data(ttl=180, show_spinner=False)
def cached_knowledge():
    repo_pairs, errors = load_repo_pairs()
    return repo_pairs, errors


try:
    repo_pairs, repo_errors = cached_knowledge()
except Exception as exc:
    repo_pairs, repo_errors = [], [str(exc)]

terms = get_terms()
pairs = get_pairs(repo_pairs)
corpus = get_corpus()

tab_translate, tab_dictionary, tab_admin, tab_about = st.tabs(
    [
        "الترجمة",
        "البحث المعجمي",
        "إدارة المنصة",
        "حول المشروع",
    ]
)


# ------------------------------------------------------------
# Translation
# ------------------------------------------------------------

with tab_translate:
    st.subheader("ترجمة النصوص")

    direction = st.selectbox(
        "اختر اتجاه الترجمة",
        list(DIRECTIONS.keys()),
        key="translation_direction",
    )

    source_lang, target_lang = DIRECTIONS[direction]

    if source_lang == "fr":
        placeholder = "Saisissez le texte à traduire..."
    elif source_lang == "ar":
        placeholder = "أدخل النص العربي المراد ترجمته..."
    else:
        placeholder = "Adɣer aḍris s tmaziɣt..."

    source_text = st.text_area(
        "النص الأصلي",
        height=200,
        placeholder=placeholder,
        key="translation_source",
    )

    translate_button = st.button(
        "ترجم النص",
        type="primary",
        use_container_width=True,
        key="translate_button",
    )

    if translate_button:
        if not clean(source_text):
            st.warning("أدخل النص أولًا.")
        else:
            with st.spinner("جارٍ البحث عن ترجمة معتمدة..."):
                exact_text, exact_origin = exact_lookup(
                    source_text,
                    direction,
                    terms,
                    pairs,
                )

            if exact_text:
                st.success(
                    f"عُثر على تطابق معتمد: {exact_origin}"
                )
                st.subheader("الترجمة")
                st.text_area(
                    "النص المترجم",
                    value=exact_text,
                    height=180,
                    key="exact_translation_result",
                )
                st.caption(
                    "هذه نتيجة مطابقة في المعرفة المعتمدة. "
                    "راجع السياق قبل استخدامها في نص صحفي."
                )

                st.session_state["last_translation"] = {
                    "source": source_text,
                    "translation": exact_text,
                    "direction": direction,
                    "origin": exact_origin,
                }

            else:
                try:
                    with st.spinner(
                        "يحلل Gemini النص ويستعين بالمصطلحات "
                        "وذاكرة الترجمة المتاحة..."
                    ):
                        result = gemini_translate(
                            source_text,
                            direction,
                            terms,
                            pairs,
                            corpus,
                        )

                    translated = result["translation"]

                    st.subheader("الترجمة المقترحة")
                    st.text_area(
                        "راجع الترجمة قبل اعتمادها",
                        value=translated,
                        height=220,
                        key="generated_translation_result",
                    )

                    st.info(
                        "الترجمة المولدة اقتراح للمراجعة، "
                        "وليست ترجمة معتمدة تلقائيًا."
                    )

                    st.session_state["last_translation"] = {
                        "source": source_text,
                        "translation": translated,
                        "direction": direction,
                        "origin": "Gemini",
                    }

                    with st.expander(
                        "عرض المصطلحات والأمثلة المستخدمة"
                    ):
                        st.markdown("**المصطلحات ذات الصلة**")
                        if result["terms"]:
                            for item in result["terms"]:
                                st.write(
                                    f"- {item.get('source', '')} "
                                    f"→ {item.get('target', '')}"
                                )
                        else:
                            st.write("لم يُعثر على مصطلحات مطابقة.")

                        st.markdown("**أمثلة ذاكرة الترجمة**")
                        if result["examples"]:
                            for item in result["examples"]:
                                st.write(
                                    f"- {item['source']} "
                                    f"→ {item['target']}"
                                )
                        else:
                            st.write("لم يُعثر على أمثلة مناسبة.")

                        if result["corpus_examples"]:
                            st.markdown(
                                "**فقرات أمازيغية مشابهة**"
                            )
                            for item in result["corpus_examples"]:
                                st.write(f"- {item['text']}")

                except Exception as exc:
                    st.error(f"تعذرت الترجمة: {exc}")

    if "last_translation" in st.session_state:
        last = st.session_state["last_translation"]

        st.divider()
        st.subheader("نسخ النتيجة")

        if st.button(
            "تجهيز نتيجة جديدة",
            key="clear_translation",
        ):
            st.session_state.pop("last_translation", None)
            st.rerun()

        st.caption(
            "يمكنك تحديد النص المترجم ونسخه. "
            "لا يُحفظ أي تصحيح تلقائيًا."
        )

        if is_admin():
            st.divider()
            st.subheader("اعتماد تصحيح الترجمة")

            corrected_text = st.text_area(
                "أدخل الصيغة النهائية التي راجعها الفريق",
                value=last["translation"],
                height=160,
                key="approved_correction_input",
            )

            if st.button(
                "اعتماد التصحيح وحفظه",
                key="approve_correction",
            ):
                try:
                    saved = save_correction(
                        source=last["source"],
                        proposed=last["translation"],
                        corrected=corrected_text,
                        direction=last["direction"],
                        origin=last["origin"],
                    )
                    st.success(
                        "تم حفظ التصحيح في قاعدة البيانات "
                        "وإضافته إلى ذاكرة الترجمة المعتمدة."
                    )
                except Exception as exc:
                    st.error(
                        f"تعذر الحفظ: {exc}"
                    )


# ------------------------------------------------------------
# Dictionary and terminology search
# ------------------------------------------------------------

with tab_dictionary:
    st.subheader("البحث في المصطلحات والمعرفة المعتمدة")

    dictionary_query = st.text_input(
        "ابحث عن مصطلح أو عبارة",
        key="dictionary_query",
    )

    if st.button(
        "بحث في المعرفة المحلية",
        key="dictionary_search",
    ):
        query = norm(dictionary_query)

        if not query:
            st.warning("أدخل مصطلحًا للبحث.")
        else:
            matches = []

            for item in terms + pairs:
                source = clean(item.get("source"))
                target = clean(item.get("target"))

                if (
                    query in norm(source)
                    or query in norm(target)
                ):
                    matches.append(item)

            if matches:
                for item in matches[:100]:
                    st.markdown(
                        f"**{item.get('source', '')}** "
                        f"↔ {item.get('target', '')}"
                    )
                    st.caption(
                        f"المصدر: {item.get('origin', 'المعرفة المحلية')} "
                        f"| الحالة: {item.get('status', 'approved')}"
                    )
            else:
                st.info("لم يُعثر على تطابق محلي.")

    st.divider()
    st.subheader("البحث في Asegzawal")

    st.caption(
        "هذا البحث اختياري ويستعلم من واجهة المعجم الخارجية؛ "
        "لا يعني ذلك استيراد المعجم كاملًا إلى Tasuqilt."
    )

    ase_query = st.text_input(
        "المصطلح المراد البحث عنه",
        key="asegzawal_query",
    )

    if st.button(
        "البحث في Asegzawal",
        key="asegzawal_search",
    ):
        if not clean(ase_query):
            st.warning("أدخل مصطلحًا للبحث.")
        else:
            with st.spinner("جارٍ البحث في المعجم..."):
                results = search_asegzawal(ase_query)

            if results:
                for item in results:
                    st.markdown(
                        f"**[{item['title']}]({item['url']})**"
                    )
                    if item["snippet"]:
                        st.write(item["snippet"])
            else:
                st.info(
                    "لم تظهر نتائج أو تعذر الاتصال بالمعجم. "
                    "يمكنك فتح موقع Asegzawal مباشرة."
                )


# ------------------------------------------------------------
# Administration
# ------------------------------------------------------------

with tab_admin:
    st.subheader("إدارة المعرفة والمصطلحات")

    if not is_admin():
        st.info(
            "تظهر أدوات الإدارة بعد تسجيل الدخول بكلمة مرور المشرف."
        )
    else:
        st.success("وضع الإدارة مفعّل.")

        if not db_ready():
            st.error(
                "قاعدة Supabase غير مهيأة. "
                "لن تعمل عمليات الحفظ الدائم حتى تضبط الإعدادات."
            )

        admin_terms, admin_pairs, admin_import, admin_corpus = st.tabs(
            [
                "إضافة مصطلح",
                "إضافة زوج ترجمة",
                "استيراد ملفات",
                "المدونة الأمازيغية",
            ]
        )

        with admin_terms:
            st.markdown("**إضافة مصطلح معياري معتمد**")

            with st.form("add_term_form"):
                term_source = st.text_input(
                    "المصطلح المصدر"
                )
                term_target = st.text_input(
                    "المصطلح الأمازيغي المعتمد"
                )

                term_direction = st.selectbox(
                    "اتجاه المصطلح",
                    list(DIRECTIONS.keys()),
                    key="term_direction",
                )

                term_note = st.text_area(
                    "ملاحظة اختيارية عن الاستعمال"
                )

                term_submit = st.form_submit_button(
                    "حفظ المصطلح المعتمد"
                )

            if term_submit:
                try:
                    sl, tl = DIRECTIONS[term_direction]
                    save_approved_term(
                        term_source,
                        term_target,
                        sl,
                        tl,
                        term_note,
                    )
                    st.success("تم حفظ المصطلح.")
                    st.cache_data.clear()
                    st.rerun()
                except Exception as exc:
                    st.error(f"تعذر حفظ المصطلح: {exc}")

        with admin_pairs:
            st.markdown(
                "**إضافة مثال ثنائي اللغة إلى ذاكرة الترجمة**"
            )

            with st.form("add_pair_form"):
                pair_source = st.text_area(
                    "النص المصدر"
                )
                pair_target = st.text_area(
                    "الترجمة المعتمدة"
                )

                pair_direction = st.selectbox(
                    "اتجاه زوج الترجمة",
                    list(DIRECTIONS.keys()),
                    key="pair_direction",
                )

                pair_origin = st.text_input(
                    "المصدر أو مرجع الترجمة",
                    value="admin",
                )

                pair_submit = st.form_submit_button(
                    "حفظ زوج الترجمة"
                )

            if pair_submit:
                try:
                    sl, tl = DIRECTIONS[pair_direction]
                    save_approved_pair(
                        pair_source,
                        pair_target,
                        sl,
                        tl,
                        pair_origin,
                    )
                    st.success(
                        "تم حفظ زوج الترجمة المعتمد."
                    )
                    st.cache_data.clear()
                    st.rerun()
                except Exception as exc:
                    st.error(f"تعذر الحفظ: {exc}")

        with admin_import:
            st.markdown("**استيراد ملفات المعرفة**")

            st.warning(
                "تُقرأ ملفات Excel وCSV من أول عمودين فقط. "
                "تحقق من عينة البيانات قبل اعتماد الاستيراد."
            )

            uploaded_files = st.file_uploader(
                "اختر ملفات المعرفة",
                type=["xlsx", "csv", "json", "txt", "md", "docx"],
                accept_multiple_files=True,
                key="knowledge_upload",
            )

            if uploaded_files:
                preview_records = []
                preview_errors = []

                for uploaded in uploaded_files:
                    try:
                        parsed = parse_file(
                            uploaded.name,
                            uploaded.getvalue(),
                        )
                        preview_records.extend(parsed)
                    except Exception as exc:
                        preview_errors.append(
                            f"{uploaded.name}: {exc}"
                        )

                st.write(
                    f"عدد أزواج الترجمة المكتشفة: "
                    f"{len(preview_records)}"
                )

                if preview_records:
                    st.dataframe(
                        pd.DataFrame(preview_records).head(100),
                        use_container_width=True,
                    )

                for error in preview_errors:
                    st.error(error)

                if st.button(
                    "تأكيد الاستيراد إلى ذاكرة الترجمة",
                    key="confirm_import",
                ):
                    try:
                        records = []

                        for item in preview_records:
                            record = dict(item)
                            record["status"] = "approved"
                            record["created_at"] = now_iso()
                            records.append(record)

                        db_upsert(
                            "translation_memory",
                            records,
                        )

                        st.success(
                            f"تم إرسال {len(records)} زوج ترجمة "
                            "إلى قاعدة البيانات."
                        )
                        st.cache_data.clear()

                    except Exception as exc:
                        st.error(
                            f"تعذر استيراد البيانات: {exc}"
                        )

        with admin_corpus:
            st.markdown(
                "**إضافة فقرات أمازيغية أحادية اللغة**"
            )

            st.caption(
                "تُخزَّن الفقرات لاسترجاع الأمثلة المشابهة. "
                "هذا الإجراء لا يدرّب Gemini ولا يغيّر أوزان النموذج."
            )

            corpus_file = st.file_uploader(
                "اختر ملف الفقرات",
                type=["txt", "md", "docx", "csv"],
                key="corpus_upload",
            )

            corpus_origin = st.text_input(
                "مصدر الفقرات",
                value="admin corpus import",
                key="corpus_origin",
            )

            if corpus_file:
                try:
                    content = corpus_file.getvalue()
                    ext = os.path.splitext(
                        corpus_file.name
                    )[1].lower()

                    paragraphs = []

                    if ext in {".txt", ".md"}:
                        paragraphs = [
                            x.strip()
                            for x in content.decode(
                                "utf-8-sig",
                                errors="replace",
                            ).split("\n")
                            if len(x.strip()) >= 20
                        ]

                    elif ext == ".docx":
                        doc = Document(io.BytesIO(content))
                        paragraphs = [
                            p.text.strip()
                            for p in doc.paragraphs
                            if len(p.text.strip()) >= 20
                        ]

                    elif ext == ".csv":
                        df = pd.read_csv(
                            io.BytesIO(content),
                            dtype=str,
                            keep_default_na=False,
                        )
                        paragraphs = [
                            str(value).strip()
                            for value in df.iloc[:, 0].tolist()
                            if len(str(value).strip()) >= 20
                        ]

                    st.write(
                        f"الفقرات المكتشفة: {len(paragraphs)}"
                    )

                    st.text_area(
                        "معاينة أولى الفقرات",
                        value="\n\n".join(paragraphs[:5]),
                        height=180,
                        disabled=True,
                    )

                    if st.button(
                        "حفظ الفقرات في المدونة",
                        key="save_corpus",
                    ):
                        count = save_corpus_paragraphs(
                            paragraphs,
                            corpus_origin,
                        )
                        st.success(
                            f"تم إرسال {count} فقرة إلى قاعدة البيانات."
                        )

                except Exception as exc:
                    st.error(
                        f"تعذرت قراءة ملف الفقرات: {exc}"
                    )


# ------------------------------------------------------------
# About and diagnostics
# ------------------------------------------------------------

with tab_about:
    st.subheader("عن Tasuqilt DZ")

    st.write(
        "تهدف Tasuqilt DZ إلى بناء منصة ترجمة أمازيغية معيارية "
        "تستند إلى المصطلحات المعتمدة وذاكرة الترجمة والأمثلة "
        "السياقية، مع الاستعانة بنموذج لغوي لتوليد ترجمات "
        "جديدة عند غياب التطابق المناسب."
    )

    st.markdown("**قواعد المصطلحات الأساسية**")

    st.markdown(
        "- Président → Aselway\n"
        "- Président de la République → Aselway n Tegduda\n"
        "- Réunion → Timlilt\n"
        "- Conseil des ministres → Aseqqamu n Yineɣlafen\n"
        "- Gouvernement → Anabaḍ\n"
        "- Alger → DZAYER TAMANEƔT"
    )

    st.warning(
        "تظل الترجمة المولدة بحاجة إلى مراجعة بشرية. "
        "ولا تعني إضافة فقرات إلى المدونة تدريب النموذج."
    )

    with st.expander("تشخيص مصادر المعرفة"):
        st.write(f"أزواج GitHub المكتشفة: {len(repo_pairs)}")
        st.write(f"الأزواج المتاحة إجمالًا: {len(pairs)}")
        st.write(f"المصطلحات المعتمدة: {len(terms)}")
        st.write(f"فقرات المدونة المحمّلة: {len(corpus)}")

        if repo_errors:
            st.warning(
                "ظهرت مشكلات عند تحميل بعض الملفات:"
            )
            for error in repo_errors[:10]:
                st.write(f"- {error}")

        if not db_ready():
            st.info(
                "لتفعيل الحفظ الدائم، أضف إعدادات Supabase "
                "إلى Streamlit Secrets."
            )

st.divider()
st.caption(
    "Tasuqilt DZ — نسخة تأسيسية قيد الاختبار. "
    "راجع الترجمات قبل نشرها رسميًا."
)
